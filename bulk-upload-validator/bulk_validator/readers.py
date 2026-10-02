"""Read CSV and XLSX upload files into the same row shape: a header list plus string cells."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


@dataclass
class Table:
    header: list[str]
    # (spreadsheet row number, cells). Row 1 is the header, so data starts at 2.
    rows: list[tuple[int, list[str]]]
    source: str


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def read_table(path: str | Path) -> Table:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        # utf-8-sig strips the BOM that spreadsheet tools often add.
        with path.open(encoding="utf-8-sig", newline="") as handle:
            raw = list(csv.reader(handle))
    elif suffix in (".xlsx", ".xlsm"):
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        raw = [[_cell_text(v) for v in row] for row in sheet.iter_rows(values_only=True)]
        workbook.close()
    else:
        raise ValueError(f"Unsupported file type {suffix!r}; use .csv or .xlsx")

    if not raw:
        return Table(header=[], rows=[], source=str(path))
    header = [h for h in raw[0]]
    width = len(header)
    rows = [(i + 2, (r + [""] * width)[:width]) for i, r in enumerate(raw[1:])]
    return Table(header=header, rows=rows, source=str(path))


def read_reference_lists(directory: str | Path) -> dict[str, set[str]]:
    """Each `<name>.csv` in the directory is a reference list: the first column holds the valid codes."""
    lists: dict[str, set[str]] = {}
    for file in sorted(Path(directory).glob("*.csv")):
        table = read_table(file)
        lists[file.stem] = {cells[0].strip() for _, cells in table.rows if cells and cells[0].strip()}
    return lists
